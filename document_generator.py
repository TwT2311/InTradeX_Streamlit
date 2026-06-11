# -*- coding: utf-8 -*-
"""
ThinkSpices Document Generator Service
Generates professional A4 print-ready PDF invoices using ReportLab
and bundled Word Docs calling python-docx. Fully modularized.
"""

import io
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm

# Formatting helpers
def format_usd(val):
    return f"${val:,.4f}" if val < 1 else f"${val:,.2f}"

def generate_pdf_bytes(doc_type, fob_data, parties, meta):
    """
    Membangun file PDF berbasis sistem A4 dengan skema warna Forest Green.
    Menggunakan ReportLab Flowables.
    """
    buffer = io.BytesIO()
    
    # Inisialisasi dokumen target dengan margin 15mm
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=15*mm,
        leftMargin=15*mm,
        topMargin=15*mm,
        bottomMargin=15*mm
    )
    
    styles = getSampleStyleSheet()
    
    # Custom Styles didasari oleh Inter/Helvetica
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        textColor=colors.white,
        leading=15,
        alignment=0 # Left
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        textColor=colors.HexColor("#C8E5C9"),
        leading=10,
        alignment=2 # Right
    )
    
    meta_label_style = ParagraphStyle(
        'MetaLabel',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        textColor=colors.HexColor("#323232"),
        leading=11,
    )
    
    meta_val_style = ParagraphStyle(
        'MetaValue',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        textColor=colors.HexColor("#4A4A4A"),
        leading=11,
    )
    
    party_header_style = ParagraphStyle(
        'PartyHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        textColor=colors.white,
        leading=10,
    )
    
    party_text_bold = ParagraphStyle(
        'PartyTextBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        textColor=colors.HexColor("#1A3A2B"),
        leading=11,
    )
    
    party_text_normal = ParagraphStyle(
        'PartyTextNormal',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        textColor=colors.HexColor("#2C2C2C"),
        leading=10,
    )
    
    table_cell_header = ParagraphStyle(
        'TableCellHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        textColor=colors.white,
        leading=10,
    )
    
    table_cell_normal = ParagraphStyle(
        'TableCellNormal',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        textColor=colors.HexColor("#2C2C2C"),
        leading=10,
    )
    
    table_cell_bold = ParagraphStyle(
        'TableCellBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        textColor=colors.HexColor("#1B4332"),
        leading=10,
    )
    
    footer_text_style = ParagraphStyle(
        'FooterText',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=7,
        textColor=colors.HexColor("#888888"),
        leading=9,
        alignment=1, # Center
    )

    story = []
    
    # ── 1. Top Header Banner Block ──
    doc_type_upper = doc_type.upper()
    title_text = "PROFORMA INVOICE" if doc_type_upper == "PI" else \
                 "COMMERCIAL INVOICE" if doc_type_upper == "CI" else \
                 "PACKING LIST" if doc_type_upper == "PL" else \
                 "SHIPPING INSTRUCTION" if doc_type_upper == "SI" else "OFFICIAL COMMERCIAL QUOTATION"
                 
    header_data = [
        [Paragraph(title_text, title_style), Paragraph("ThinkSpices Trade Intelligence Partner", subtitle_style)]
    ]
    # Total printable width is of A4 (210mm - 30mm margins) = 180mm
    header_table = Table(header_data, colWidths=[110*mm, 70*mm])
    header_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#1B4332")),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3*mm),
        ('TOPPADDING', (0,0), (-1,-1), 3*mm),
        ('LEFTPADDING', (0,0), (0,0), 6*mm),
        ('RIGHTPADDING', (-1,-1), (-1,-1), 6*mm),
    ]))
    story.append(header_table)
    story.append(Spacer(1, 4*mm))
    
    # ── 2. Information Metadata Grid block ──
    meta_left_label_text = "Document No:"
    meta_left_val_text = meta.get("docNumber", "—")
    meta_right_label_text = "Issue Date:"
    meta_right_val_text = meta.get("issueDate", "—")
    
    if doc_type in ["quote", "pi"]:
        alt_label = "Validity Days:"
        alt_val = f"{meta.get('validityDays', 14)} days"
        sec_label = "Incoterm:"
        sec_val = f"FOB {fob_data.get('loading_port', 'Origin Port')}"
    else:
        alt_label = meta.get("additionalFieldLabel", "Ref PIN:")
        alt_val = meta.get("additionalFieldValue", "—")
        sec_label = meta.get("secondaryFieldLabel", "Incoterm Basis:")
        sec_val = meta.get("secondaryFieldValue", "—")
        
    meta_data = [
        [
            Paragraph(meta_left_label_text, meta_label_style), Paragraph(meta_left_val_text, meta_val_style),
            Paragraph(meta_right_label_text, meta_label_style), Paragraph(meta_right_val_text, meta_val_style)
        ],
        [
            Paragraph(alt_label, meta_label_style), Paragraph(alt_val, meta_val_style),
            Paragraph(sec_label, meta_label_style), Paragraph(sec_val, meta_val_style)
        ]
    ]
    meta_table = Table(meta_data, colWidths=[30*mm, 60*mm, 30*mm, 60*mm])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#F0FAF4")),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#D8F3DC")),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 2*mm),
        ('BOTTOMPADDING', (0,0), (-1,-1), 2*mm),
        ('LEFTPADDING', (0,0), (-1,-1), 4*mm),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 5*mm))
    
    # ── 3. Parties side-by-side Layout blocks ──
    col_w = 87*mm
    parties_header_data = [
        [Paragraph("SELLER / EXPORTER", party_header_style), "", Paragraph("BUYER / CONSIGNEE", party_header_style)]
    ]
    parties_header_table = Table(parties_header_data, colWidths=[col_w, 6*mm, col_w])
    parties_header_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (0,0), colors.HexColor("#2D6A4F")),
        ('BACKGROUND', (2,0), (2,0), colors.HexColor("#2D6A4F")),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 1.5*mm),
        ('BOTTOMPADDING', (0,0), (-1,-1), 1.5*mm),
        ('LEFTPADDING', (0,0), (-1,-1), 4*mm),
    ]))
    story.append(parties_header_table)
    
    seller_body = [
        Paragraph(parties.get("sellerName", "—"), party_text_bold),
        Paragraph(parties.get("sellerAddress", "—").replace("\n", "<br/>"), party_text_normal),
        Paragraph(f"Country: {parties.get('sellerCountry', 'Indonesia')}", party_text_normal),
        Paragraph(f"Phone: {parties.get('sellerPhone', '—')}", party_text_normal),
        Paragraph(f"Email: {parties.get('sellerEmail', '—')}", party_text_normal),
    ]
    buyer_body = [
        Paragraph(parties.get("buyerName", "—"), party_text_bold),
        Paragraph(parties.get("buyerAddress", "—").replace("\n", "<br/>"), party_text_normal),
        Paragraph(f"Country: {parties.get('buyerCountry', '—')}", party_text_normal),
        Paragraph(f"Phone: {parties.get('buyerPhone', '—')}", party_text_normal),
        Paragraph(f"Email: {parties.get('buyerEmail', '—')}", party_text_normal),
    ]
    
    parties_content_data = [
        [seller_body, "", buyer_body]
    ]
    parties_content_table = Table(parties_content_data, colWidths=[col_w, 6*mm, col_w])
    parties_content_table.setStyle(TableStyle([
        ('BOX', (0,0), (0,0), 0.5, colors.HexColor("#DCDCDC")),
        ('BOX', (2,0), (2,0), 0.5, colors.HexColor("#DCDCDC")),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('TOPPADDING', (0,0), (-1,-1), 2.5*mm),
        ('BOTTOMPADDING', (0,0), (-1,-1), 2.5*mm),
        ('LEFTPADDING', (0,0), (-1,-1), 4*mm),
        ('RIGHTPADDING', (0,0), (-1,-1), 4*mm),
    ]))
    story.append(parties_content_table)
    story.append(Spacer(1, 5*mm))
    
    # ── 4. Goods specifications Matrix ──
    commodity = fob_data.get("commodity", "—")
    hs_code = fob_data.get("hs_code", "—")
    volume = fob_data.get("volume_kg", 0)
    price_per_kg = fob_data.get("fob_price_per_kg", 0.0)
    total_val = fob_data.get("fob_total_usd", 0.0)
    
    goods_headers = [
        Paragraph("Description of Goods", table_cell_header),
        Paragraph("HS Code", table_cell_header),
        Paragraph("Qty (Kg)", table_cell_header),
        Paragraph("Unit Price (USD)", table_cell_header),
        Paragraph("Amount (USD)", table_cell_header)
    ]
    
    goods_rows = [
        Paragraph(commodity, table_cell_bold),
        Paragraph(hs_code, table_cell_normal),
        Paragraph(f"{volume:,.0f} Kg", table_cell_normal),
        Paragraph(format_usd(price_per_kg), table_cell_normal),
        Paragraph(format_usd(total_val), table_cell_normal)
    ]
    
    goods_data = [
        goods_headers,
        goods_rows
    ]
    
    goods_table = Table(goods_data, colWidths=[65*mm, 25*mm, 25*mm, 35*mm, 30*mm])
    goods_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#2D6A4F")),
        ('BACKGROUND', (0,1), (-1,1), colors.HexColor("#F8F8F8")),
        ('BOX', (0,1), (-1,1), 0.5, colors.HexColor("#E0E0E0")),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 2*mm),
        ('BOTTOMPADDING', (0,0), (-1,-1), 2*mm),
        ('LEFTPADDING', (0,0), (-1,-1), 3*mm),
        ('RIGHTPADDING', (-1,0), (-1,-1), 3*mm),
    ]))
    story.append(goods_table)
    
    # Total Row block
    total_port_name = f"TOTAL FOB PORT VALUE ({fob_data.get('loading_port', '—')}):"
    total_row_data = [
        [Paragraph(total_port_name, table_cell_bold), Paragraph(format_usd(total_val), table_cell_bold)]
    ]
    total_table = Table(total_row_data, colWidths=[150*mm, 30*mm])
    total_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#D8F3DC")),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#2D6A4F")),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 1.5*mm),
        ('BOTTOMPADDING', (0,0), (-1,-1), 1.5*mm),
        ('LEFTPADDING', (0,0), (0,0), 3*mm),
        ('RIGHTPADDING', (-1,-1), (-1,-1), 3*mm),
    ]))
    story.append(total_table)
    story.append(Spacer(1, 4*mm))
    
    # ── 5. Packaging Weight Specs ──
    specs_data = [
        [
            Paragraph("<b>Packaging:</b>", table_cell_normal),
            Paragraph(f"{fob_data.get('total_units_needed', 0)} units of standard packaged crop", table_cell_normal),
            Paragraph("<b>Net Weight:</b>", table_cell_normal),
            Paragraph(f"{volume:,.0f} Kg", table_cell_normal),
            Paragraph("<b>Gross Weight:</b>", table_cell_normal),
            Paragraph(f"{fob_data.get('gross_weight_kg', 0):,.1f} Kg", table_cell_normal)
        ]
    ]
    specs_table = Table(specs_data, colWidths=[20*mm, 60*mm, 20*mm, 30*mm, 25*mm, 25*mm])
    specs_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#FCFCFC")),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#DCDCDC")),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 2*mm),
        ('BOTTOMPADDING', (0,0), (-1,-1), 2*mm),
        ('LEFTPADDING', (0,0), (-1,-1), 3*mm),
    ]))
    story.append(specs_table)
    story.append(Spacer(1, 4*mm))
    
    # ── 6. Remittance Banking Blocks ──
    terms_body = meta.get("paymentTerms", "—")
    bank_body = meta.get("bankDetails", "—")
    
    remit_header_table = Table([
        [Paragraph("PAYMENT TERMS AGREED", party_header_style), "", Paragraph("REMITTANCE BANK DETAILS", party_header_style)]
    ], colWidths=[col_w, 6*mm, col_w])
    remit_header_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (0,0), colors.HexColor("#2D6A4F")),
        ('BACKGROUND', (2,0), (2,0), colors.HexColor("#2D6A4F")),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 1.5*mm),
        ('BOTTOMPADDING', (0,0), (-1,-1), 1.5*mm),
        ('LEFTPADDING', (0,0), (-1,-1), 3*mm),
    ]))
    story.append(remit_header_table)
    
    remit_content_table = Table([
        [Paragraph(terms_body.replace('\n', '<br/>'), table_cell_normal), "", Paragraph(bank_body.replace('\n', '<br/>'), table_cell_normal)]
    ], colWidths=[col_w, 6*mm, col_w])
    remit_content_table.setStyle(TableStyle([
        ('BOX', (0,0), (0,0), 0.5, colors.HexColor("#DCDCDC")),
        ('BOX', (2,0), (2,0), 0.5, colors.HexColor("#DCDCDC")),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('TOPPADDING', (0,0), (-1,-1), 2*mm),
        ('BOTTOMPADDING', (0,0), (-1,-1), 2*mm),
        ('LEFTPADDING', (0,0), (-1,-1), 3*mm),
        ('RIGHTPADDING', (0,0), (-1,-1), 3*mm),
    ]))
    story.append(remit_content_table)
    story.append(Spacer(1, 4*mm))
    
    # ── 7. Regulatory Rules & Compliances ──
    qa_header = Paragraph("SPECIAL INSTRUCTIONS & GLOBAL COMPLIANCE RULES", ParagraphStyle('QaH', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=8, textColor=colors.HexColor("#1B4332")))
    qa_line = Table([[""]], colWidths=[180*mm])
    qa_line.setStyle(TableStyle([
        ('LINEBELOW', (0,0), (-1,-1), 1, colors.HexColor("#2D6A4F")),
        ('BOTTOMPADDING', (0,0), (-1,-1), 1),
        ('TOPPADDING', (0,0), (-1,-1), 1),
    ]))
    
    notes_formatted = meta.get("notes", "—").replace("\n", "<br/> &nbsp; ")
    notes_paragraph = Paragraph(notes_formatted, table_cell_normal)
    
    story.append(qa_header)
    story.append(qa_line)
    story.append(Spacer(1, 1.5*mm))
    story.append(notes_paragraph)
    story.append(Spacer(1, 8*mm))
    
    # ── 8. Corporate Signatures ──
    sig_cols_data = [
        [
            Paragraph("_____________________________<br/><b>AUTHORISED SELLER SIGNATURE</b><br/>PT Magastu Indoprime Group", table_cell_normal),
            "",
            Paragraph("_____________________________<br/><b>BUYER ACCEPTANCE STAMP</b><br/>" + parties.get('buyerName', 'Consignee Purchaser'), table_cell_normal)
        ]
    ]
    sig_table = Table(sig_cols_data, colWidths=[80*mm, 20*mm, 80*mm])
    sig_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('TOPPADDING', (0,0), (-1,-1), 2*mm),
        ('ALIGN', (2,0), (2,-1), 'RIGHT'),
    ]))
    
    story.append(KeepTogether([sig_table]))
    story.append(Spacer(1, 8*mm))
    
    # Subtle Footer branding centring
    story.append(Paragraph("Generated securely via ThinkSpices Premium Trade Desk.", footer_text_style))
    
    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def generate_docx_bytes(doc_type, fob_data, parties, meta):
    """
    Membangun file Word (.docx) berbasis python-docx.
    Menghasilkan bytes yang dapat diunduh di Streamlit.
    """
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import parse_xml, OxmlElement
    from docx.oxml.ns import nsdecls, qn
    
    doc = Document()
    
    # Set Margin standard 0.6 inci
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)
        
    doc_type_upper = doc_type.upper()
    title_text = "PROFORMA INVOICE" if doc_type_upper == "PI" else \
                 "COMMERCIAL INVOICE" if doc_type_upper == "CI" else \
                 "PACKING LIST" if doc_type_upper == "PL" else \
                 "SHIPPING INSTRUCTION" if doc_type_upper == "SI" else "OFFICIAL EXPORT QUOTATION"
                 
    # Custom XML helpers for styling word cells
    def set_cell_background(cell, color_hex):
        shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
        cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))
        
    def add_cell_border(cell, **kwargs):
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = OxmlElement('w:tcBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            edge_data = kwargs.get(edge)
            if edge_data:
                tag = 'w:{}'.format(edge)
                element = OxmlElement(tag)
                element.set(qn('w:val'), edge_data.get('val', 'single'))
                element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
                element.set(qn('w:space'), '0')
                element.set(qn('w:color'), edge_data.get('color', 'auto'))
                tcBorders.append(element)
        tcPr.append(tcBorders)

    # Title header
    p_title = doc.add_paragraph()
    r_title = p_title.add_run(title_text)
    r_title.bold = True
    r_title.size = Pt(18)
    r_title.font.name = 'Helvetica'
    r_title.font.color.rgb = RGBColor(27, 67, 50) # Forest Green
    
    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("THINKSPICES - EXPEDITED TRADE PROCUREMENT SYSTEM")
    r_sub.bold = True
    r_sub.size = Pt(8.5)
    r_sub.font.name = 'Helvetica'
    r_sub.font.color.rgb = RGBColor(45, 106, 79)
    
    # Metadata Table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.autofit = False
    
    meta_left_ref = meta_table.cell(0, 0)
    meta_right_date = meta_table.cell(0, 1)
    meta_left_alt = meta_table.cell(1, 0)
    meta_right_sec = meta_table.cell(1, 1)
    
    for row in meta_table.rows:
        row.cells[0].width = Inches(3.25)
        row.cells[1].width = Inches(3.25)
        
    p1 = meta_left_ref.paragraphs[0]
    p1.add_run("Document No: ").bold = True
    p1.add_run(meta.get("docNumber", "—"))
    p1.style.font.size = Pt(8)
    
    p2 = meta_right_date.paragraphs[0]
    p2.add_run("Issue Date: ").bold = True
    p2.add_run(meta.get("issueDate", "—"))
    p2.style.font.size = Pt(8)
    
    p3 = meta_left_alt.paragraphs[0]
    if doc_type in ["quote", "pi"]:
        p3.add_run("Validity Period: ").bold = True
        p3.add_run(f"{meta.get('validityDays', 14)} days")
    else:
        p3.add_run(f"{meta.get('additionalFieldLabel', 'Ref')}: ").bold = True
        p3.add_run(meta.get("additionalFieldValue", "—"))
    p3.style.font.size = Pt(8)
        
    p4 = meta_right_sec.paragraphs[0]
    if doc_type in ["quote", "pi"]:
        p4.add_run("Incoterm Standard: ").bold = True
        p4.add_run(f"FOB {fob_data.get('loading_port', '—')}, Indonesia")
    else:
        p4.add_run(f"{meta.get('secondaryFieldLabel', 'Incoterm')}: ").bold = True
        p4.add_run(meta.get("secondaryFieldValue", "—"))
    p4.style.font.size = Pt(8)
        
    for row in meta_table.rows:
        for cell in row.cells:
            set_cell_background(cell, "EBF8EB")
            add_cell_border(cell, top={'sz': 2, 'val': 'single', 'color': 'D8F3DC'},
                                  bottom={'sz': 2, 'val': 'single', 'color': 'D8F3DC'},
                                  left={'sz': 2, 'val': 'single', 'color': 'D8F3DC'},
                                  right={'sz': 2, 'val': 'single', 'color': 'D8F3DC'})
            
    doc.add_paragraph() # Spacer
    
    # Contracting parties table
    p_parties = doc.add_paragraph()
    r_parties = p_parties.add_run("CONTRACTING PARTIES / TRANSACTION INVOLVEMENT")
    r_parties.bold = True
    r_parties.size = Pt(10)
    r_parties.font.color.rgb = RGBColor(27, 67, 50)
    
    parties_table = doc.add_table(rows=1, cols=2)
    parties_table.autofit = False
    
    cell_seller = parties_table.cell(0, 0)
    cell_buyer = parties_table.cell(0, 1)
    cell_seller.width = Inches(3.25)
    cell_buyer.width = Inches(3.25)
    
    # Exporter content
    ps1 = cell_seller.paragraphs[0]
    rs1_head = ps1.add_run("SELLER / EXPORTER\n")
    rs1_head.bold = True
    rs1_head.font.color.rgb = RGBColor(45, 106, 79)
    ps1.add_run(parties.get("sellerName", "—")).bold = True
    ps1.add_run(f"\n{parties.get('sellerAddress', '—')}\nCountry: {parties.get('sellerCountry', 'Indonesia')}\nPhone: {parties.get('sellerPhone', '—')}\nEmail: {parties.get('sellerEmail', '—')}")
    ps1.style.font.size = Pt(8)
    
    # Buyer content
    pb1 = cell_buyer.paragraphs[0]
    rb1_head = pb1.add_run("BUYER / CONSIGNEE\n")
    rb1_head.bold = True
    rb1_head.font.color.rgb = RGBColor(45, 106, 79)
    pb1.add_run(parties.get("buyerName", "—")).bold = True
    pb1.add_run(f"\n{parties.get('buyerAddress', '—')}\nCountry: {parties.get('buyerCountry', '—')}\nPhone: {parties.get('buyerPhone', '—')}\nEmail: {parties.get('buyerEmail', '—')}")
    pb1.style.font.size = Pt(8)
    
    for cell in (cell_seller, cell_buyer):
        set_cell_background(cell, "F0F5F0")
        add_cell_border(cell, top={'sz': 4, 'val': 'single', 'color': 'DCDCDC'},
                              bottom={'sz': 4, 'val': 'single', 'color': 'DCDCDC'},
                              left={'sz': 4, 'val': 'single', 'color': 'DCDCDC'},
                              right={'sz': 4, 'val': 'single', 'color': 'DCDCDC'})
                              
    doc.add_paragraph() # Spacer
    
    # Goods Specifications Header
    goods_head = doc.add_paragraph()
    r_ghead = goods_head.add_run("SPECIFICATIONS & VALUE BREAKDOWN")
    r_ghead.bold = True
    r_ghead.size = Pt(10)
    r_ghead.font.color.rgb = RGBColor(27, 67, 50)
    
    goods_table = doc.add_table(rows=3, cols=5)
    goods_table.autofit = False
    
    col_widths = [Inches(2.5), Inches(1.0), Inches(0.8), Inches(1.1), Inches(1.1)]
    for row in goods_table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width
            
    header_titles = ["Spice Commodity / Grade", "HS Code", "Net Qty", "Price / Kg (USD)", "Total Amount"]
    for idx, title in enumerate(header_titles):
        cell = goods_table.cell(0, idx)
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        p.style.font.size = Pt(8)
        set_cell_background(cell, "2D6A4F")
        
    goods_data_list = [
        fob_data.get("commodity", "—"),
        fob_data.get("hs_code", "—"),
        f"{fob_data.get('volume_kg', 0):,.0f} Kg",
        format_usd(fob_data.get("fob_price_per_kg", 0.0)),
        format_usd(fob_data.get("fob_total_usd", 0.0))
    ]
    for idx, data_text in enumerate(goods_data_list):
        cell = goods_table.cell(1, idx)
        p = cell.paragraphs[0]
        r = p.add_run(data_text)
        if idx == 0:
            r.bold = True
        p.style.font.size = Pt(8)
        set_cell_background(cell, "F9FBF9")
        add_cell_border(cell, top={'sz': 2, 'val': 'single', 'color': 'E2E8F0'},
                              bottom={'sz': 2, 'val': 'single', 'color': 'E2E8F0'},
                              left={'sz': 2, 'val': 'single', 'color': 'E2E8F0'},
                              right={'sz': 2, 'val': 'single', 'color': 'E2E8F0'})
                              
    total_cell_label = goods_table.cell(2, 0)
    total_cell_value = goods_table.cell(2, 4)
    total_cell_label.merge(goods_table.cell(2, 1)).merge(goods_table.cell(2, 2)).merge(goods_table.cell(2, 3))
    
    p_tot_label = total_cell_label.paragraphs[0]
    r_ptl = p_tot_label.add_run(f"TOTAL FOB PORT VALUE (EX-WORKS INDONESIA + FEES) ({fob_data.get('loading_port', '—')}):")
    r_ptl.bold = True
    r_ptl.font.color.rgb = RGBColor(27, 67, 50)
    p_tot_label.style.font.size = Pt(7.5)
    
    p_tot_val = total_cell_value.paragraphs[0]
    r_ptv = p_tot_val.add_run(format_usd(fob_data.get("fob_total_usd", 0.0)))
    r_ptv.bold = True
    r_ptv.font.color.rgb = RGBColor(27, 67, 50)
    p_tot_val.style.font.size = Pt(8)
    
    for cell in (total_cell_label, total_cell_value):
        set_cell_background(cell, "D8F3DC")
        add_cell_border(cell, top={'sz': 4, 'val': 'single', 'color': '2D6A4F'},
                              bottom={'sz': 4, 'val': 'single', 'color': '2D6A4F'},
                              left={'sz': 4, 'val': 'single', 'color': '2D6A4F'},
                              right={'sz': 4, 'val': 'single', 'color': '2D6A4F'})
                              
    p_specs = doc.add_paragraph()
    r_specs = p_specs.add_run(f"Logistics Specs: {fob_data.get('total_units_needed', 0)} units | Net: {fob_data.get('volume_kg', 0):,.0f} Kg | Est Gross: {fob_data.get('gross_weight_kg', 0):,.1f} Kg")
    r_specs.italic = True
    r_specs.size = Pt(8)
    
    p_payment_head = doc.add_paragraph()
    r_ph = p_payment_head.add_run("TERMS & REMITTANCE SCHEME")
    r_ph.bold = True
    r_ph.size = Pt(10)
    r_ph.font.color.rgb = RGBColor(27, 67, 50)
    
    # Remittance Bank Table
    remit_table = doc.add_table(rows=1, cols=2)
    remit_table.autofit = False
    
    c_pay = remit_table.cell(0, 0)
    c_bank = remit_table.cell(0, 1)
    c_pay.width = Inches(3.25)
    c_bank.width = Inches(3.25)
    
    p_pay = c_pay.paragraphs[0]
    p_pay.add_run("TERMS OF PAYMENT\n").bold = True
    p_pay.add_run(meta.get("paymentTerms", "—"))
    p_pay.style.font.size = Pt(7.5)
    
    p_bank = c_bank.paragraphs[0]
    p_bank.add_run("REMITTANCE BANK DETAILS\n").bold = True
    p_bank.add_run(meta.get("bankDetails", "—"))
    p_bank.style.font.size = Pt(7.5)
    
    for cell in (c_pay, c_bank):
        set_cell_background(cell, "F4FAF4")
        add_cell_border(cell, top={'sz': 2, 'val': 'single', 'color': 'E2E8F0'},
                              bottom={'sz': 2, 'val': 'single', 'color': 'E2E8F0'},
                              left={'sz': 2, 'val': 'single', 'color': 'E2E8F0'},
                              right={'sz': 2, 'val': 'single', 'color': 'E2E8F0'})
                              
    doc.add_paragraph() # Spacer
    
    p_inst_head = doc.add_paragraph()
    r_ihead = p_inst_head.add_run("SPECIAL INSTRUCTIONS & REGULATORY COMPLIANCE")
    r_ihead.bold = True
    r_ihead.size = Pt(10)
    r_ihead.font.color.rgb = RGBColor(27, 67, 50)
    
    p_inst_body = doc.add_paragraph()
    r_ibody = p_inst_body.add_run(meta.get("notes", "—"))
    r_ibody.size = Pt(7.5)
    
    doc.add_paragraph() # Spacer
    
    # Signatures
    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.autofit = False
    
    c_sig_sel = sig_table.cell(0, 0)
    c_sig_buy = sig_table.cell(0, 1)
    
    c_sig_sel.width = Inches(3.25)
    c_sig_buy.width = Inches(3.25)
    
    p_s1 = c_sig_sel.paragraphs[0]
    p_s1.add_run("\n\nAUTHORISED SELLER SIGNATURE\nPT Magastu Indoprime Group").bold = True
    p_s1.style.font.size = Pt(7.5)
    
    p_s2 = c_sig_buy.paragraphs[0]
    p_s2.add_run(f"\n\nBUYER TRANSACTION ACCEPTANCE\n{parties.get('buyerName', 'Consignee')}").bold = True
    p_s2.style.font.size = Pt(7.5)
    
    for cell in (c_sig_sel, c_sig_buy):
        add_cell_border(cell, top={'sz': 2, 'val': 'single', 'color': 'E0E0E0'})
        
    doc.add_paragraph() # Spacer
    p_footer = doc.add_paragraph()
    p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_foot = p_footer.add_run("Generated securely via ThinkSpices Premium Trade Desk.")
    r_foot.italic = True
    r_foot.size = Pt(7)
    r_foot.font.color.rgb = RGBColor(120, 120, 120)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output.getvalue()
